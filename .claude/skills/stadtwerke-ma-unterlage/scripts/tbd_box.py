#!/usr/bin/env python3
"""Gelbe TBD-Kaesten: anlegen, auflisten, entfernen.

Ueber allem, was der Skill nicht fertig machen kann, liegt ein gelber Kasten mit
fortlaufender Nummer, damit der Kollege beim Durchblaettern sofort sieht, was zu tun ist.
Drei Sorten, gleiche Farbe, unterschiedliche erste Zeile:

    TC       think-cell aktualisieren, Excel-Block <Name>
    LOGO     Logo oder Icon ersetzen
    PRUEFEN  Pruefen (Urteil des Nutzers)

    python3 tbd_box.py deck.pptx --add befehle.json --out deck_v2.pptx
    python3 tbd_box.py deck.pptx --list                    Liste = TBD-Liste des Reports
    python3 tbd_box.py deck.pptx --remove --out final.pptx

Format der Befehlsdatei (Folie, Sorte, Text, optional Position/Groesse in cm; ohne Position
legt sich der Kasten ueber das genannte Shape oder mittig auf die Folie):

    [
      {"slide": 15, "sorte": "TC", "text": "Excel-Block GuV-Chart", "ueber": "id:120"},
      {"slide": 14, "sorte": "TC", "text": "Excel-Block Umsatzring", "ueber": "name:Chart 57"},
      {"slide": 13, "sorte": "LOGO", "text": "Logo Stadtwerke Netz GmbH", "x": 3.0, "y": 8.0, "w": 4.0, "h": 1.5},
      {"slide": 18, "sorte": "PRUEFEN", "text": "zwei fallspezifische Ziele bestaetigen"}
    ]

Die Kaesten heissen TBD-<Nr> und werden von inspect_deck.py --vergleich als bewusste
Ergaenzung erkannt, nicht als fremdes Textfeld. --remove entfernt genau diese Shapes und
sonst nichts.

Benoetigt: pip install python-pptx
"""
import argparse
import json
import os
import re
import sys

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Cm, Pt
except ImportError:
    sys.exit("Fehlt: pip install python-pptx")

for strom in (sys.stdout, sys.stderr):
    try:
        strom.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

SORTEN = {
    "TC": "think-cell aktualisieren",
    "LOGO": "Logo oder Icon ersetzen",
    "PRUEFEN": "Prüfen",
}
NAME = re.compile(r"^TBD-(\d+)$")
GELB = RGBColor(0xFF, 0xE6, 0x00)
SCHWARZ = RGBColor(0x00, 0x00, 0x00)


def alle_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from alle_shapes(sh.shapes)


def vorhandene(prs):
    out = []
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            m = NAME.match(sh.name or "")
            if m:
                out.append((int(m.group(1)), i, sh))
    return sorted(out, key=lambda x: x[0])


def finde(slide, adresse):
    art, _, wert = adresse.partition(":")
    for sh in alle_shapes(slide.shapes):
        if art == "id" and str(sh.shape_id) == wert:
            return sh
        if art == "name" and sh.name == wert:
            return sh
        if art == "ph" and sh.is_placeholder and str(sh.placeholder_format.idx) == wert:
            return sh
    return None


def lege_an(prs, slide, nr, sorte, text, x=None, y=None, w=None, h=None, ueber=None):
    if ueber:
        ziel = finde(slide, ueber)
        if ziel is None or ziel.left is None:
            raise ValueError(f"'{ueber}' auf Folie nicht gefunden oder ohne Position")
        left, top, width, height = ziel.left, ziel.top, ziel.width, ziel.height
        # Kasten etwas kleiner als das Objekt, damit die Umgebung sichtbar bleibt
        if height > Cm(1.6):
            top = top + int(height * 0.3)
            height = int(height * 0.4)
    else:
        width = Cm(w) if w else Cm(8)
        height = Cm(h) if h else Cm(1.2)
        left = Cm(x) if x is not None else int((prs.slide_width - width) / 2)
        top = Cm(y) if y is not None else int((prs.slide_height - height) / 2)
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.name = f"TBD-{nr}"
    box.fill.solid()
    box.fill.fore_color.rgb = GELB
    box.line.color.rgb = SCHWARZ
    box.line.width = Pt(0.75)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Cm(0.15)
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = f"TBD {nr} | {sorte} | {SORTEN[sorte]}"
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = SCHWARZ
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = text
    r2.font.size = Pt(8)
    r2.font.color.rgb = SCHWARZ
    return box


def liste(prs):
    zeilen = vorhandene(prs)
    if not zeilen:
        print("Keine TBD-Kaesten in der Datei.")
        return 0
    print("| Nr | Folie | Sorte | Was zu tun ist |")
    print("|---|---|---|---|")
    for nr, folie, sh in zeilen:
        t = sh.text_frame.text.split("\n")
        kopf = t[0].split("|")
        sorte = kopf[1].strip() if len(kopf) > 1 else ""
        rest = " ".join(x.strip() for x in t[1:])
        print(f"| {nr} | {folie} | {sorte} | {rest} |")
    print(f"\n{len(zeilen)} Kasten/Kaesten. Die Nummern sind die der TBD-Liste im Report.")
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("--add", metavar="BEFEHLE.json")
    ap.add_argument("--list", action="store_true")
    ap.add_argument("--remove", action="store_true")
    ap.add_argument("--out")
    a = ap.parse_args()
    prs = Presentation(a.pptx)

    if a.list:
        sys.exit(liste(prs))

    if a.out and os.path.abspath(a.out) == os.path.abspath(a.pptx):
        sys.exit("Ziel- und Quelldatei sind identisch - anderen Namen fuer --out waehlen.")

    if a.remove:
        if not a.out:
            sys.exit("--remove braucht --out")
        n = 0
        for nr, folie, sh in vorhandene(prs):
            sh._element.getparent().remove(sh._element)
            n += 1
        prs.save(a.out)
        print(f"{n} TBD-Kasten/Kaesten entfernt. Gespeichert: {a.out}")
        return

    if a.add:
        if not a.out:
            sys.exit("--add braucht --out")
        with open(a.add, encoding="utf-8") as f:
            befehle = json.load(f)
        nr = max([n for n, _, _ in vorhandene(prs)] or [0])
        fehler = []
        for b in befehle:
            try:
                folie = int(b["slide"])
                sorte = b["sorte"].upper()
                if sorte not in SORTEN:
                    raise ValueError(f"Sorte '{b['sorte']}' unbekannt (TC, LOGO, PRUEFEN)")
                if not 1 <= folie <= len(prs.slides):
                    raise ValueError(f"Folie {folie} gibt es nicht")
                nr += 1
                lege_an(prs, prs.slides[folie - 1], nr, sorte, b.get("text", ""),
                        b.get("x"), b.get("y"), b.get("w"), b.get("h"), b.get("ueber"))
                print(f"TBD {nr:>2}  F{folie:>2}  {sorte:<8} {b.get('text', '')[:60]}")
            except (KeyError, ValueError) as e:
                fehler.append(f"{b}: {e}")
        if fehler:
            print("\nNicht angelegt:", file=sys.stderr)
            for f in fehler:
                print(f"  {f}", file=sys.stderr)
            sys.exit(1)
        prs.save(a.out)
        print(f"\nGespeichert: {a.out}. Liste mit: tbd_box.py {a.out} --list")
        return

    ap.print_help()


if __name__ == "__main__":
    main()
