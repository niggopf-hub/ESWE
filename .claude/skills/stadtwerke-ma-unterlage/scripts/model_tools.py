#!/usr/bin/env python3
"""Werkzeuge fuer das Financial Model im Hausformat.

    python3 model_tools.py rename vorlage.xlsx --kuerzel ESWE --name "ESWE Versorgungs AG" --out modell.xlsx
    python3 model_tools.py check  modell.xlsx            Kontrollzeilen ueber LibreOffice durchrechnen
    python3 model_tools.py blocks modell.xlsx            think-cell-Bloecke (Overview FS ab Spalte J) als Markdown
    python3 model_tools.py euro   modell.xlsx            Euro-Zeichen, T EUR-Reste, Prozent-Woerter finden

rename ersetzt das Platzhalter-Kuerzel GES in Blattnamen, Formeln und Texten sowie
<Gesellschaft> im Cover und im Quellenblatt. check schreibt eine durchgerechnete Kopie
(soffice --headless) und liest jede Zeile, deren Beschriftung mit "Kontrolle" beginnt;
alles, was nicht null ist, wird gemeldet. Fehlerwerte (#REF!, #NAME?, #DIV/0!) in
irgendeiner Zelle sind ebenfalls ein Befund.

Benoetigt: pip install openpyxl; fuer check zusaetzlich LibreOffice (soffice im PATH).
"""
import argparse
import os
import shutil
import subprocess
import sys
import tempfile

try:
    import openpyxl
except ImportError:
    sys.exit("Fehlt: pip install openpyxl")

for strom in (sys.stdout, sys.stderr):
    try:
        strom.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass


# ---------------------------------------------------------------- rename

def rename(pfad, kuerzel, name, out):
    if os.path.abspath(pfad) == os.path.abspath(out):
        sys.exit("Ziel- und Quelldatei sind identisch. Die Vorlage wird nicht ueberschrieben.")
    wb = openpyxl.load_workbook(pfad)
    alt = [ws.title for ws in wb.worksheets if ws.title.startswith("GES ")]
    for ws in wb.worksheets:
        if ws.title.startswith("GES "):
            ws.title = f"{kuerzel} " + ws.title[4:]
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    v = c.value
                    v2 = v.replace("'GES ", f"'{kuerzel} ").replace("‹Gesellschaft›", name)
                    v2 = v2.replace("GES ", f"{kuerzel} ") if v2.startswith("GES ") else v2
                    if v2 != v:
                        c.value = v2
                        n += 1
    wb.save(out)
    print(f"Blaetter umbenannt: {', '.join(alt)} -> Kuerzel '{kuerzel}'")
    print(f"{n} Zellen angepasst. Gespeichert: {out}")
    print("Naechster Schritt: Cover!C24 (erstes Geschaeftsjahr), KPIs!C45 (Einheit) und KPIs!C57 (Net-Debt-Stufe) setzen.")


# ---------------------------------------------------------------- check

def durchrechnen(pfad):
    """Laesst LibreOffice die Mappe neu berechnen und gibt den Pfad der Kopie zurueck."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None, "LibreOffice (soffice) nicht gefunden - Kontrollen koennen nicht durchgerechnet werden."
    tmp = tempfile.mkdtemp(prefix="modell_")
    cmd = [soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, pfad]
    try:
        subprocess.run(cmd, check=True, capture_output=True, timeout=180)
    except Exception as e:  # noqa: BLE001
        return None, f"LibreOffice-Lauf fehlgeschlagen: {e}"
    ziel = os.path.join(tmp, os.path.basename(pfad))
    if not ziel.lower().endswith(".xlsx"):
        ziel = os.path.splitext(ziel)[0] + ".xlsx"
    if not os.path.exists(ziel):
        return None, "LibreOffice hat keine Ausgabedatei erzeugt."
    return ziel, None


FEHLER = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "Err:")


def check(pfad):
    ziel, err = durchrechnen(pfad)
    if err:
        print(f"{err}\nRechne stattdessen mit pycel (pip install pycel).")
        return check_pycel(pfad)
    wb = openpyxl.load_workbook(ziel, data_only=True)
    befunde = []
    kontrollen = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and any(v.startswith(f) for f in FEHLER):
                    befunde.append((ws.title, c.coordinate, f"Fehlerwert {v}"))
        for r in range(1, ws.max_row + 1):
            label = ws.cell(r, 2).value or ws.cell(r, 1).value
            if not isinstance(label, str) or not label.lower().startswith("kontrolle"):
                continue
            kontrollen += 1
            for col in range(3, 8):
                v = ws.cell(r, col).value
                if v is None or v == "n. v." or v == "":
                    continue
                try:
                    if abs(float(v)) > 0.0005:
                        befunde.append((ws.title, ws.cell(r, col).coordinate,
                                        f"{label[:60]} = {v}"))
                except (TypeError, ValueError):
                    befunde.append((ws.title, ws.cell(r, col).coordinate, f"{label[:60]} = {v!r}"))
    print(f"Durchgerechnete Kopie: {ziel}")
    print(f"{kontrollen} Kontrollzeilen gelesen.")
    if not befunde:
        print("Alle Kontrollen auf null, keine Fehlerwerte.")
        print("Ungeprueft bleibt: ob die Eingaben die Zahlen des Abschlusses sind. Stichprobe gegen das PDF.")
        return 0
    print(f"{len(befunde)} Befund(e):")
    for blatt, zelle, was in befunde:
        print(f"  {blatt}!{zelle}: {was}")
    return 1


def check_pycel(pfad):
    """Kontrollzeilen ohne LibreOffice: pycel wertet die Formeln in Python aus."""
    try:
        from pycel import ExcelCompiler
    except ImportError:
        print("pycel fehlt: pip install pycel")
        return 2
    xl = ExcelCompiler(filename=pfad)
    wb = openpyxl.load_workbook(pfad)
    befunde, kontrollen = [], 0
    for ws in wb.worksheets:
        for r in range(1, ws.max_row + 1):
            label = ws.cell(r, 2).value or ws.cell(r, 1).value
            if not isinstance(label, str) or not label.lower().startswith("kontrolle"):
                continue
            kontrollen += 1
            for col in range(3, 8):
                adr = f"'{ws.title}'!{ws.cell(r, col).coordinate}"
                if ws.cell(r, col).value is None:
                    continue
                try:
                    v = xl.evaluate(adr)
                except Exception as e:  # noqa: BLE001
                    befunde.append((ws.title, ws.cell(r, col).coordinate, f"nicht auswertbar: {e}"))
                    continue
                if v in (None, "", "n. v."):
                    continue
                try:
                    if abs(float(v)) > 0.0005:
                        befunde.append((ws.title, ws.cell(r, col).coordinate, f"{label[:60]} = {v}"))
                except (TypeError, ValueError):
                    if str(v).startswith("#"):
                        befunde.append((ws.title, ws.cell(r, col).coordinate, f"Fehlerwert {v}"))
    print(f"{kontrollen} Kontrollzeilen mit pycel gerechnet.")
    if not befunde:
        print("Alle Kontrollen auf null.")
        print("Ungeprueft bleibt: ob die Eingaben die Zahlen des Abschlusses sind. Stichprobe gegen das PDF.")
        return 0
    print(f"{len(befunde)} Befund(e):")
    for blatt, zelle, was in befunde:
        print(f"  {blatt}!{zelle}: {was}")
    return 1


# ---------------------------------------------------------------- blocks

def blocks(pfad):
    ziel, err = durchrechnen(pfad)
    wb = openpyxl.load_workbook(ziel or pfad, data_only=bool(ziel))
    xl = None
    if err:
        try:
            from pycel import ExcelCompiler
            xl = ExcelCompiler(filename=pfad)
        except ImportError:
            print(f"Hinweis: {err} Es werden Formeln statt Werte gezeigt.\n")
    ws = wb["Overview FS"]

    def wert(c):
        v = c.value
        if xl is not None and isinstance(v, str) and v.startswith("="):
            try:
                return xl.evaluate(f"'Overview FS'!{c.coordinate}")
            except Exception:  # noqa: BLE001
                return v
        return v
    r = 1
    out = []
    while r <= ws.max_row:
        kopf = ws.cell(r, 10).value
        fill = ws.cell(r, 10).fill.fgColor.rgb if ws.cell(r, 10).fill and ws.cell(r, 10).fill.fgColor else None
        if isinstance(kopf, str) and fill and str(fill).endswith("003D7C"):
            spalten = [wert(ws.cell(r, c)) for c in range(11, 17)]
            spalten = [s for s in spalten if s is not None]
            out.append(f"\n**{kopf}**\n")
            out.append("| | " + " | ".join(fmt(s) for s in spalten) + " |")
            out.append("|---|" + "---|" * len(spalten))
            r += 1
            while r <= ws.max_row and ws.cell(r, 10).value is not None:
                zeile = [fmt(wert(ws.cell(r, c)), ws.cell(r, c).number_format)
                         for c in range(11, 11 + len(spalten))]
                out.append(f"| {wert(ws.cell(r, 10))} | " + " | ".join(zeile) + " |")
                r += 1
        else:
            r += 1
    print("\n".join(out) if out else "Keine think-cell-Bloecke ab Spalte J gefunden.")
    return 0


def fmt(v, numfmt=""):
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        if "%" in numfmt:
            return f"{v * 100:.1f} %".replace(".", ",")
        if "x" in numfmt:
            return f"{v:.1f}x".replace(".", ",")
        if abs(v) >= 1000 and float(v).is_integer():
            return f"{v:,.0f}".replace(",", ".")
        if float(v).is_integer() and "0.0" not in numfmt:
            return f"{v:.0f}"
        return f"{v:,.1f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return str(v)


# ---------------------------------------------------------------- euro

def euro(pfad):
    wb = openpyxl.load_workbook(pfad)
    befunde = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and not v.startswith("="):
                    if "€" in v:
                        befunde.append((ws.title, c.coordinate, "Euro-Zeichen", v[:60]))
                    elif "T EUR" in v or "TEUR" in v:
                        befunde.append((ws.title, c.coordinate, "T EUR statt Mio. EUR", v[:60]))
                    elif "Prozent" in v:
                        befunde.append((ws.title, c.coordinate, "'Prozent' ausgeschrieben", v[:60]))
                if "€" in (c.number_format or ""):
                    befunde.append((ws.title, c.coordinate, "Euro-Zeichen im Zahlenformat", c.number_format))
    if not befunde:
        print("Keine Euro-Zeichen, keine T EUR, kein 'Prozent'.")
        return 0
    print(f"{len(befunde)} Befund(e):")
    for blatt, zelle, was, t in befunde:
        print(f"  {blatt}!{zelle}: {was}  ->  {t!r}")
    return 1


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("rename")
    r.add_argument("xlsx")
    r.add_argument("--kuerzel", required=True, help="z. B. ESWE, DVV, SWK")
    r.add_argument("--name", required=True, help="voller Name der Gesellschaft")
    r.add_argument("--out", required=True)
    for cmd in ("check", "blocks", "euro"):
        s = sub.add_parser(cmd)
        s.add_argument("xlsx")
    a = p.parse_args()
    if a.cmd == "rename":
        rename(a.xlsx, a.kuerzel, a.name, a.out)
        return
    sys.exit({"check": check, "blocks": blocks, "euro": euro}[a.cmd](a.xlsx))


if __name__ == "__main__":
    main()
