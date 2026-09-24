#!/usr/bin/env python3
"""Folien aus einer anderen Datei uebernehmen, ohne Master zu mischen.

    python3 import_slide.py ziel.pptx --von quelle.pptx --folien 1,2 --nach 17 --out ziel_v2.pptx
    python3 import_slide.py ziel.pptx --von alt_blau.pptx --folien 20,21 --nach 17 --umfaerben --out ziel_v2.pptx

Fuer jede Quellfolie entsteht im Ziel eine neue Folie mit dem Layout "Nur Titel" (bzw. dem
ersten Inhaltslayout), und die Shapes der Quellfolie werden als XML hineinkopiert. Der
Master des Ziels bleibt der einzige; Theme-Farben und Schriften der Platzhalter stellen
sich damit von selbst um. Bilder, Diagramme und OLE-Objekte werden nicht kopiert und
einzeln gemeldet - sie brauchen Beziehungen, die beim Kopieren nicht mitkommen.

--umfaerben ersetzt hart gesetzte Farben des blauen Templates durch die des gruenen. Das
ist fuer Textkaesten und Rechtecke ausreichend, fuer Diagramme nicht.

--nach N fuegt die Folien nach Folie N ein (0 = an den Anfang); ohne --nach ans Ende.

Benoetigt: pip install python-pptx
"""
import argparse
import copy
import os
import re
import sys

try:
    from pptx import Presentation
except ImportError:
    sys.exit("Fehlt: pip install python-pptx")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import vorlage  # noqa: E402


def rollenboxen_ergaenzen(prs, neu):
    """Quelle und Fussnote sind im gruenen Template freie Textfelder, keine klonbaren
    Platzhalter. Fehlen sie auf der neuen Folie, werden sie von einer Folie desselben
    Decks kopiert, die sie hat - nicht neu gezogen."""
    ergaenzt = []
    for rolle in ("quelle", "fussnote"):
        if vorlage.finde_rolle(neu, rolle):
            continue
        for s in prs.slides:
            if s is neu:
                continue
            treffer = [sh for sh in vorlage.finde_rolle(s, rolle) if not sh.is_placeholder]
            if treffer:
                neu.shapes._spTree.append(copy.deepcopy(treffer[0]._element))
                ergaenzt.append(rolle)
                break
    return ergaenzt

FARBEN = {  # blau -> gruen
    "003D7C": "44501A", "3157E3": "8FB1A4", "5B8DEF": "A6ABCF", "A5CAFB": "B0BE94",
    "D1E4FD": "D1E0D7", "EAF2FE": "D1E0D7", "FF9900": "DFB36C", "17365D": "44501A",
}
NICHT = ("pic", "graphicFrame", "oleObj")


def inhaltslayout(prs):
    for m in prs.slide_masters:
        for l in m.slide_layouts:
            if l.name == "Nur Titel":
                return l
    for m in prs.slide_masters:
        for l in m.slide_layouts:
            if "inhalt" in l.name.lower() or "titel" in l.name.lower():
                return l
    return prs.slide_layouts[0]


def tag(el):
    return el.tag.split("}")[-1]


def umfaerben(el):
    n = 0
    for c in el.iter():
        if tag(c) == "srgbClr":
            v = (c.get("val") or "").upper()
            if v in FARBEN:
                c.set("val", FARBEN[v])
                n += 1
    return n


def verschiebe(prs, slide, nach):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    el = ids[-1]
    lst.remove(el)
    lst.insert(nach, el)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("pptx")
    ap.add_argument("--von", required=True)
    ap.add_argument("--folien", required=True, help="Foliennummern der Quelle, kommagetrennt")
    ap.add_argument("--nach", type=int, help="nach dieser Folie im Ziel einfuegen")
    ap.add_argument("--umfaerben", action="store_true")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    if os.path.abspath(a.out) == os.path.abspath(a.pptx):
        sys.exit("Ziel- und Quelldatei sind identisch - anderen Namen fuer --out waehlen.")

    ziel = Presentation(a.pptx)
    quelle = Presentation(a.von)
    layout = inhaltslayout(ziel)
    nummern = [int(x) for x in a.folien.split(",")]
    einfuegen = a.nach
    for nr in nummern:
        if not 1 <= nr <= len(quelle.slides):
            sys.exit(f"Quellfolie {nr} gibt es nicht (1..{len(quelle.slides)})")
        q = quelle.slides[nr - 1]
        neu = ziel.slides.add_slide(layout)
        # Leere Inhaltsplatzhalter des Layouts entfernen; Titel, Kolumne, Fussnote, Quelle bleiben
        for ph in list(neu.placeholders):
            if ph.placeholder_format.idx not in (0, 13, 11, 20, 2, 18):
                ph._element.getparent().remove(ph._element)
        n_ok, uebersprungen, farben = 0, [], 0
        for sh in q.shapes:
            el = sh._element
            if sh.is_placeholder:
                # Platzhalter der Quelle nicht mitnehmen: ihre Indizes bedeuten im Ziel etwas
                # anderes (ph:14 war Quelle, ist Subline). Nur den Titeltext uebernehmen.
                if sh.placeholder_format.idx == 0 and sh.has_text_frame and neu.shapes.title is not None:
                    neu.shapes.title.text_frame.text = sh.text_frame.text
                elif sh.has_text_frame and sh.text_frame.text.strip():
                    uebersprungen.append(f"{sh.name} (Platzhalter, Text: {sh.text_frame.text.strip()[:40]!r})")
                continue
            if tag(el) in NICHT or el.xpath(".//*[local-name()='blip' or local-name()='chart' or local-name()='oleObj']"):
                uebersprungen.append(sh.name)
                continue
            kopie = copy.deepcopy(el)
            if a.umfaerben:
                farben += umfaerben(kopie)
            neu.shapes._spTree.append(kopie)
            n_ok += 1
        boxen = rollenboxen_ergaenzen(ziel, neu)
        if einfuegen is not None:
            verschiebe(ziel, neu, einfuegen)
            einfuegen += 1
        print(f"Quellfolie {nr}: {n_ok} Shapes uebernommen"
              + (f", {'/'.join(boxen)}-Box von einer Nachbarfolie kopiert" if boxen else "")
              + (f", {farben} Farben umgestellt" if a.umfaerben else "")
              + (f"; NICHT kopiert (Bild/Diagramm/OLE/Platzhalter): {', '.join(uebersprungen)}" if uebersprungen else ""))
    ziel.save(a.out)
    print(f"Gespeichert: {a.out}. Danach: inspect_deck.py --slide N, Titel/Kolumne/Quelle mit fill_deck.py setzen, Sichtpruefung.")


if __name__ == "__main__":
    main()
